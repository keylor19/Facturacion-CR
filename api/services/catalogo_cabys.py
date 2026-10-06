"""
Catálogo CABYS local.

El catálogo completo (~20 000 códigos) lo publica el BCCR como archivo de Excel;
se importa a la tabla `cabys` y las búsquedas se hacen aquí, sin consultar a
Hacienda (más rápido, sin límites de consumo y funciona si Hacienda está caída).
Mientras no se haya importado, se van guardando los códigos que devuelven las
consultas a Hacienda.

Actualizar cuando el BCCR publique una versión nueva: cambie CABYS_URL y ejecute
`python -m api.cli cargar-cabys` (o --archivo con el Excel descargado).
"""
import hashlib
import io
import logging
import re
import time
import unicodedata
from decimal import Decimal

import requests
from sqlalchemy import and_, case, delete, func, or_, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import SQLAlchemyError

from api.models.database import CodigoCabys, SessionLocal, utcnow
from config.settings import get_settings

logger = logging.getLogger(__name__)

MINIMO_COMPLETO = 10_000      # con esta cantidad de códigos se considera cargado el catálogo oficial
_completo = {"valor": False, "hasta": 0.0}


class CatalogoError(Exception):
    pass


def normalizar(texto: str) -> str:
    """Minúsculas y sin tildes (para buscar "cafe" y encontrar "Café")."""
    sin_tildes = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in sin_tildes if not unicodedata.combining(c)).lower().strip()


def _porcentaje(valor) -> Decimal:
    """El Excel trae el IVA como fracción (0.13), "13%" o "Exento"; Hacienda como porcentaje (13)."""
    texto = str(valor).strip().lower()
    if texto.startswith("exent"):
        return Decimal("0.00")
    if "%" in texto:
        return Decimal(texto.replace("%", "").strip()).quantize(Decimal("0.01"))
    n = Decimal(texto)
    return (n * 100 if n <= 1 else n).quantize(Decimal("0.01"))


def _numero(valor: Decimal):
    return int(valor) if valor == valor.to_integral_value() else float(valor)


def _item(c: CodigoCabys) -> dict:
    """Mismo formato que devuelve el API de Hacienda."""
    return {"codigo": c.codigo, "descripcion": c.descripcion, "categorias": c.categorias or [],
            "impuesto": _numero(Decimal(c.impuesto)), "uri": "", "estado": "", "origen": "catalogo_local"}


def completo() -> bool:
    """¿Está importado el catálogo oficial? (se recalcula cada 10 minutos)."""
    if time.monotonic() < _completo["hasta"]:
        return _completo["valor"]
    try:
        with SessionLocal() as s:
            valor = (s.scalar(select(func.count()).select_from(CodigoCabys)) or 0) >= MINIMO_COMPLETO
    except SQLAlchemyError:
        valor = False
    _completo.update(valor=valor, hasta=time.monotonic() + 600)
    return valor


PALABRAS_VACIAS = {"de", "del", "la", "las", "el", "los", "y", "o", "para", "con", "sin", "en", "por", "a", "al",
                   "un", "una", "uso", "tipo"}


def _raiz(palabra: str) -> str:
    """
    Comienzo de la palabra sin la terminación, para que coincidan plural, singular y
    derivados: computadoras ~ computadores, contables ~ contabilidad, comercial ~ comerciales.
    """
    return palabra if len(palabra) <= 5 else palabra[:max(5, len(palabra) - 3)]


def buscar(texto: str, top: int = 20) -> dict | None:
    """
    Busca por palabras en la descripción y las categorías (sin tildes, plural o
    singular). Primero los que tienen todas las palabras; si ninguno, los que
    tienen más. Ordena por coincidencia exacta y descripción más corta. None si falla la BD.
    """
    palabras = [p for p in re.findall(r"[a-z0-9ñ]+", normalizar(texto)) if len(p) >= 2 and p not in PALABRAS_VACIAS][:6]
    if not palabras:
        return {"total": 0, "cantidad": 0, "cabys": []}
    texto_total = CodigoCabys.busqueda + " " + CodigoCabys.busqueda_categorias
    coincide = [texto_total.op("~")(r"\m" + _raiz(p)) for p in palabras]
    # 3 puntos si la palabra exacta está en la descripción, 2 si empieza igual, 1 si solo está en las categorías
    puntaje = sum(case((CodigoCabys.busqueda.op("~")(rf"\m{p}\M"), 3),
                       (CodigoCabys.busqueda.op("~")(r"\m" + _raiz(p)), 2),
                       (c, 1), else_=0) for p, c in zip(palabras, coincide))
    try:
        with SessionLocal() as s:
            filtro, parcial = and_(*coincide), False
            total = s.scalar(select(func.count()).select_from(CodigoCabys).where(filtro)) or 0
            if total == 0 and len(palabras) > 1:
                filtro, parcial = or_(*coincide), True
                total = s.scalar(select(func.count()).select_from(CodigoCabys).where(filtro)) or 0
            filas = s.scalars(select(CodigoCabys).where(filtro)
                              .order_by(puntaje.desc(), func.length(CodigoCabys.descripcion)).limit(top)).all()
            lista = [_item(c) for c in filas]
    except SQLAlchemyError as exc:
        logger.warning("No se pudo buscar en el catálogo CABYS local: %s", exc)
        return None
    return {"total": total, "cantidad": len(lista), "cabys": lista, "parcial": parcial}


def por_codigo(codigo: str) -> list | None:
    """[item] si existe, [] si no, None si falla la BD."""
    try:
        with SessionLocal() as s:
            c = s.get(CodigoCabys, codigo)
            return [_item(c)] if c else []
    except SQLAlchemyError as exc:
        logger.warning("No se pudo leer el catálogo CABYS local: %s", exc)
        return None


def _valores(codigo: str, descripcion: str, impuesto, categorias: list) -> dict:
    categorias = [str(x).strip() for x in categorias if x]
    return {"codigo": codigo, "descripcion": descripcion.strip(), "impuesto": _porcentaje(impuesto),
            "categorias": categorias, "busqueda": normalizar(descripcion),
            "busqueda_categorias": normalizar(" ".join(categorias)), "actualizado_en": utcnow()}


def _upsert(s, filas: list[dict]) -> None:
    if not filas:
        return
    stmt = pg_insert(CodigoCabys).values(filas)
    s.execute(stmt.on_conflict_do_update(index_elements=["codigo"], set_={
        k: stmt.excluded[k] for k in ("descripcion", "impuesto", "categorias", "busqueda", "busqueda_categorias", "actualizado_en")}))


def guardar_resultados(lista) -> None:
    """Guarda los códigos que devolvió Hacienda (mientras no esté importado el catálogo)."""
    filas = []
    for x in lista or []:
        try:
            if len(str(x.get("codigo", ""))) == 13 and x.get("descripcion") and x.get("impuesto") is not None:
                filas.append(_valores(str(x["codigo"]), x["descripcion"], Decimal(str(x["impuesto"])) / 100,
                                      x.get("categorias") or []))
        except (ArithmeticError, ValueError, AttributeError):
            continue
    if not filas:
        return
    try:
        with SessionLocal() as s:
            _upsert(s, list({f["codigo"]: f for f in filas}.values()))
            s.commit()
    except SQLAlchemyError as exc:
        logger.warning("No se pudieron guardar códigos CABYS: %s", exc)


def importar_excel(contenido: bytes) -> dict:
    """
    Importa el Excel oficial del BCCR (hoja "Catálogo": columnas Categoría 1..9 con
    su descripción e Impuesto). Reemplaza el catálogo: borra los códigos que ya no vienen.
    """
    import openpyxl

    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # archivo dañado o que no es Excel
        raise CatalogoError(f"No se pudo leer el archivo de Excel: {exc}") from exc

    hoja, cols, inicio = None, None, None
    for ws in libro.worksheets:
        for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
            titulos = [normalizar(str(v)) if v is not None else "" for v in fila]
            if "categoria 9" in titulos and "impuesto" in titulos:
                hoja, inicio = ws, i + 1
                cols = {"codigo": titulos.index("categoria 9"), "impuesto": titulos.index("impuesto"),
                        "categorias": [titulos.index(f"descripcion (categoria {n})") for n in range(1, 9)
                                       if f"descripcion (categoria {n})" in titulos]}
                break
        if hoja:
            break
    if not hoja:
        raise CatalogoError("El archivo no tiene la hoja del catálogo CABYS (columnas 'Categoría 9' e 'Impuesto')")

    filas, codigos, omitidas = {}, set(), 0
    for fila in hoja.iter_rows(min_row=inicio, values_only=True):
        codigo = str(fila[cols["codigo"]] or "").strip()
        descripcion = fila[cols["codigo"] + 1]
        if len(codigo) != 13 or not codigo.isdigit() or not descripcion or fila[cols["impuesto"]] is None:
            omitidas += 1
            continue
        try:
            filas[codigo] = _valores(codigo, str(descripcion), fila[cols["impuesto"]],
                                     [fila[c] for c in cols["categorias"]])
        except (ArithmeticError, ValueError):
            omitidas += 1
    if not filas:
        raise CatalogoError("El archivo no trae códigos CABYS válidos")

    valores = list(filas.values())
    with SessionLocal() as s:
        for i in range(0, len(valores), 1000):
            _upsert(s, valores[i:i + 1000])
        eliminados = 0
        if len(valores) >= MINIMO_COMPLETO:   # solo si es el catálogo completo
            eliminados = s.execute(delete(CodigoCabys).where(CodigoCabys.codigo.notin_(list(filas)))).rowcount
        s.commit()
    _completo["hasta"] = 0.0
    resultado = {"importados": len(valores), "eliminados": eliminados, "omitidas": omitidas}
    logger.info("Catálogo CABYS importado: %s", resultado)
    return resultado


def descargar_e_importar(url: str | None = None) -> dict:
    url = url or get_settings().CABYS_URL
    try:
        resp = requests.get(url, timeout=120, headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
    except requests.RequestException as exc:
        raise CatalogoError(f"No se pudo descargar el catálogo CABYS de {url}: {exc}") from exc
    return importar_excel(resp.content)


# --- Actualización automática ---------------------------------------------------
# El BCCR publica cada versión nueva del catálogo en una noticia con el enlace al
# Excel (la dirección cambia en cada versión). Todos los días se revisan las
# noticias recientes y el archivo configurado; si hay un archivo nuevo o cambió,
# se importa solo y queda en la bitácora.
BCCR = "https://www.bccr.fi.cr"
NOTICIAS_BCCR = BCCR + "/cr/es/noticias.html"
_NAVEGADOR = {"User-Agent": "Mozilla/5.0"}


def _estado() -> dict:
    from api.models.database import RegistroHacienda
    try:
        with SessionLocal() as s:
            r = s.scalar(select(RegistroHacienda).where(RegistroHacienda.tipo == "sistema",
                                                        RegistroHacienda.clave == "cabys"))
            return dict(r.datos or {}) if r else {}
    except SQLAlchemyError:
        return {}


def _guardar_estado(datos: dict) -> None:
    from api.models.database import RegistroHacienda, gen_uuid
    ahora = utcnow()
    valores = {"ruta": "bccr", "parametros": {}, "encontrado": True, "datos": datos,
               "actualizado_en": ahora, "usado_en": ahora}
    with SessionLocal() as s:
        s.execute(pg_insert(RegistroHacienda).values(id=gen_uuid(), tipo="sistema", clave="cabys", **valores)
                  .on_conflict_do_update(constraint="uq_registro_hacienda", set_=valores))
        s.commit()


def descubrir_versiones() -> list[str]:
    """Excel del catálogo publicados en las noticias recientes del BCCR (el más reciente primero)."""
    try:
        portada = requests.get(NOTICIAS_BCCR, timeout=30, headers=_NAVEGADOR).text
    except requests.RequestException as exc:
        logger.warning("No se pudieron revisar las noticias del BCCR: %s", exc)
        return []
    noticias = dict.fromkeys(re.findall(
        r'href="(/cr/es/noticias/listado-de-noticias/\d{4}/[^"]*(?:cabys|catalogo-de-bienes)[^"]*\.html)"', portada, re.I))
    archivos = []
    for ruta in list(noticias)[:5]:
        try:
            pagina = requests.get(BCCR + ruta, timeout=30, headers=_NAVEGADOR).text
        except requests.RequestException:
            continue
        for href in re.findall(r'href="([^"]*catalogo-de-bienes[^"]*\.xlsx)"', pagina, re.I):
            if "equivalencia" not in href.lower():
                archivos.append(href if href.startswith("http") else BCCR + href)
    return sorted(dict.fromkeys(archivos), key=_fecha_archivo, reverse=True)


def _fecha_archivo(url: str | None) -> str:
    """Los archivos del BCCR empiezan con la fecha de publicación (2025-04-01-...)."""
    m = re.search(r"(\d{4}-\d{2}-\d{2})", url or "")
    return m.group(1) if m else ""


def _firma(url: str) -> str | None:
    """ETag / fecha / tamaño del archivo, para saber si cambió sin descargarlo."""
    try:
        r = requests.head(url, timeout=30, headers=_NAVEGADOR, allow_redirects=True)
        if r.status_code != 200:
            return None
        return "|".join(r.headers.get(h, "") for h in ("ETag", "Last-Modified", "Content-Length"))
    except requests.RequestException:
        return None


def actualizar_si_hay_version_nueva() -> dict:
    """Importa el catálogo si hay una versión nueva publicada, si cambió el archivo o si aún no se ha cargado."""
    estado = _estado()
    configurado = get_settings().CABYS_URL
    descubiertos = descubrir_versiones()
    # Una versión descubierta en las noticias reemplaza a la configurada si es más reciente
    actual = estado.get("url") or configurado
    candidatos = [u for u in descubiertos[:1] if _fecha_archivo(u) > _fecha_archivo(actual)] + [actual]
    for url in dict.fromkeys(candidatos):
        firma = _firma(url)
        if url == estado.get("url") and firma and firma == estado.get("firma") and completo():
            return {"cambios": False, "url": url}
        try:
            resp = requests.get(url, timeout=120, headers=_NAVEGADOR)
            resp.raise_for_status()
        except requests.RequestException as exc:
            logger.warning("No se pudo descargar el catálogo CABYS de %s: %s", url, exc)
            continue
        huella = hashlib.sha256(resp.content).hexdigest()
        if url == estado.get("url") and huella == estado.get("sha256") and completo():
            _guardar_estado({**estado, "firma": firma})
            return {"cambios": False, "url": url}
        resultado = importar_excel(resp.content)
        _guardar_estado({"url": url, "firma": firma, "sha256": huella, "importado_en": utcnow().isoformat(),
                         **resultado})
        _registrar_en_bitacora(url, resultado, anterior=estado.get("url"))
        return {"cambios": True, "url": url, **resultado}
    return {"cambios": False, "error": "No se pudo descargar ningún catálogo"}


def _registrar_en_bitacora(url: str, resultado: dict, anterior: str | None) -> None:
    from api.services import auditoria
    detalle = (f"{'Versión nueva' if anterior and anterior != url else 'Catálogo'} importado de {url}: "
               f"{resultado['importados']} códigos, {resultado['eliminados']} eliminados")
    with SessionLocal() as s:
        auditoria.registrar(s, "cabys.actualizado", actor="Sistema (actualización automática)", detalle=detalle)
